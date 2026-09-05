import io
import json
import pytest
import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from main import app
from consistency.schemas import (
    NormalizedSource,
    SourceType,
    UCKR,
    DocumentMetadata,
    FactItem,
    EntityItem,
    ClaimItem,
    StatisticItem,
    RelationshipItem,
    OutputGenerationConfig
)
from consistency.extractor import (
    extract_from_raw_text,
    extract_from_pdf_bytes,
    extract_from_docx_bytes,
    extract_from_txt_bytes,
    extract_source
)
from consistency.analyzer import analyze_source_and_build_uckr, create_fallback_uckr
from consistency.fact_registry import FactRegistry, get_or_create_registry, get_registry
from consistency.generators import (
    SummaryGenerator,
    LinkedInGenerator,
    PresentationGenerator,
    VideoGenerator,
    TwitterGenerator,
    AdvisoryGenerator
)
from consistency.engine import ConsistencyEngine

client = TestClient(app)


def create_sample_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(text: str) -> bytes:
    doc = Document()
    doc.add_heading("Smart Cities AI Architecture", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SAMPLE_ARTICLE_TEXT = """
Artificial intelligence is accelerating urban transformation worldwide.
AI optimizes traffic management and dynamic signal routing across major metropolitan corridors, reducing congestion by 35%.
In energy grids, machine learning improves power efficiency and cuts carbon emissions by 28%.
Urban planners utilize predictive AI simulations to forecast population growth and infrastructure demand through 2035.
"""


# ---------------------------------------------------------------------------
# Stage 1 & 2: Source Extraction Layer Tests
# ---------------------------------------------------------------------------

def test_extract_raw_text():
    normalized = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    assert normalized.source_id.startswith("SRC-")
    assert normalized.title == "AI in Smart Cities"
    assert normalized.source_type == SourceType.RAW_TEXT
    assert len(normalized.raw_text) > 50
    assert len(normalized.sections) >= 1
    assert "traffic management" in normalized.raw_text


def test_extract_pdf():
    pdf_bytes = create_sample_pdf("Artificial Intelligence can optimize traffic management and energy efficiency.")
    normalized = extract_from_pdf_bytes(pdf_bytes, filename="smart_cities.pdf")
    assert normalized.source_id.startswith("SRC-")
    assert normalized.source_type == SourceType.PDF
    assert "traffic management" in normalized.raw_text
    assert normalized.metadata["filename"] == "smart_cities.pdf"


def test_extract_docx():
    docx_bytes = create_sample_docx("AI can improve energy efficiency across municipal grids.")
    normalized = extract_from_docx_bytes(docx_bytes, filename="smart_cities.docx")
    assert normalized.source_id.startswith("SRC-")
    assert normalized.source_type == SourceType.DOCX
    assert "energy efficiency" in normalized.raw_text
    assert normalized.metadata["filename"] == "smart_cities.docx"


def test_extract_txt():
    txt_bytes = SAMPLE_ARTICLE_TEXT.encode("utf-8")
    normalized = extract_from_txt_bytes(txt_bytes, filename="notes.txt")
    assert normalized.source_id.startswith("SRC-")
    assert normalized.source_type == SourceType.TXT
    assert "predictive AI" in normalized.raw_text


def test_extract_source_dispatcher():
    # PDF dispatch
    pdf_bytes = create_sample_pdf("Dispatch PDF test text.")
    norm_pdf = extract_source(file_bytes=pdf_bytes, filename="doc.pdf")
    assert norm_pdf.source_type == SourceType.PDF

    # Raw text dispatch
    norm_text = extract_source(raw_text="Dispatch text test.")
    assert norm_text.source_type == SourceType.RAW_TEXT


# ---------------------------------------------------------------------------
# Stage 3 & 4: Content Understanding & UCKR Creation Tests
# ---------------------------------------------------------------------------

def test_fallback_uckr_construction():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="Smart Cities AI")
    uckr = create_fallback_uckr(norm)

    assert uckr.document.id == norm.source_id
    assert uckr.document.title == "Smart Cities AI"
    assert len(uckr.facts) >= 2
    assert uckr.facts[0].id == "F001"
    assert uckr.facts[1].id == "F002"
    assert uckr.facts[0].importance > 0.0
    assert len(uckr.statistics) >= 1  # Should catch 35% or 28%
    assert len(uckr.key_concepts) >= 1


def test_analyze_source_and_build_uckr(monkeypatch):
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="Smart Cities AI")

    fake_uckr_json = {
        "document": {
            "id": norm.source_id,
            "title": "Smart Cities AI",
            "domain": "Artificial Intelligence"
        },
        "core_topic": "Artificial Intelligence in Smart Cities",
        "summary": "AI is used to improve urban traffic and energy efficiency.",
        "facts": [
            {
                "id": "F001",
                "statement": "AI can optimize traffic management.",
                "importance": 0.95,
                "source_reference": "page_1",
                "confidence": 0.94,
                "category": "Traffic"
            },
            {
                "id": "F002",
                "statement": "AI can improve energy efficiency.",
                "importance": 0.88,
                "source_reference": "page_2",
                "confidence": 0.91,
                "category": "Energy"
            }
        ],
        "entities": [
            {"id": "E001", "name": "Artificial Intelligence", "type": "Technology"},
            {"id": "E002", "name": "Smart Cities", "type": "Domain"}
        ],
        "claims": [
            {"id": "C001", "claim": "AI improves traffic management.", "support": "section_1", "status": "Supported"}
        ],
        "statistics": [
            {"id": "S001", "value": "35%", "context": "energy and traffic reduction", "source_reference": "sec_1"}
        ],
        "key_concepts": ["AI", "Traffic Optimization", "Energy Efficiency"],
        "relationships": [
            {"source": "Artificial Intelligence", "relation": "improves", "target": "Traffic Management"}
        ]
    }

    def fake_qwen(prompt):
        return f"```json\n{json.dumps(fake_uckr_json)}\n```"

    monkeypatch.setattr("consistency.analyzer.generate_with_qwen", fake_qwen)

    uckr = analyze_source_and_build_uckr(norm)
    assert uckr.core_topic == "Artificial Intelligence in Smart Cities"
    assert len(uckr.facts) == 2
    assert uckr.facts[0].id == "F001"
    assert uckr.facts[0].statement == "AI can optimize traffic management."
    assert uckr.facts[1].id == "F002"
    assert len(uckr.entities) == 2
    assert uckr.entities[0].name == "Artificial Intelligence"
    assert len(uckr.statistics) == 1
    assert uckr.statistics[0].value == "35%"


# ---------------------------------------------------------------------------
# Stage 5 & 6: Fact Registry & Audit Tests
# ---------------------------------------------------------------------------

def test_fact_registry_indexing_and_search():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="Smart Cities")
    uckr = create_fallback_uckr(norm)
    registry = FactRegistry.from_uckr(uckr)

    assert len(registry.list_facts()) == len(uckr.facts)
    fact_f001 = registry.get_fact("F001")
    assert fact_f001 is not None
    assert fact_f001.id == "F001"

    # Test search
    search_res = registry.search_facts("traffic")
    assert len(search_res) >= 1

    # Test critical facts
    critical = registry.get_critical_facts(min_importance=0.85)
    assert len(critical) >= 1
    assert all(f.importance >= 0.85 for f in critical)

    # Test validation
    val_res = registry.validate_fact_ids(["F001", "F002", "F999"])
    assert "F001" in val_res["valid"]
    assert "F999" in val_res["invalid"]
    assert val_res["is_valid"] is False


def test_fact_registry_audit_matrix():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="Smart Cities")
    uckr = create_fallback_uckr(norm)
    registry = FactRegistry.from_uckr(uckr)

    sample_outputs = {
        "summary": {
            "text": "Executive summary paragraph.",
            "source_facts": ["F001", "F002"]
        },
        "linkedin": {
            "post_content": "LinkedIn post.",
            "source_facts": ["F001"]
        },
        "presentation": {
            "slides": [
                {"slide_number": 1, "source_facts": ["F001"]},
                {"slide_number": 2, "source_facts": ["F002", "F003"]}
            ],
            "source_facts": ["F001", "F002", "F003"]
        },
        "video": {
            "storyboard": [
                {"scene_number": 1, "source_facts": ["F001"]}
            ],
            "source_facts": ["F001"]
        }
    }

    audit = registry.audit_deliverables(sample_outputs)
    assert audit.total_registered_facts == len(uckr.facts)
    assert audit.total_cited_facts >= 3
    assert audit.fact_coverage_percentage > 0.0
    assert audit.citation_integrity_score == 1.0
    assert "summary" in audit.traceability_matrix["F001"]
    assert "linkedin" in audit.traceability_matrix["F001"]
    assert "video" in audit.traceability_matrix["F001"]


# ---------------------------------------------------------------------------
# Stage 7: Grounded Deliverable Generators Tests
# ---------------------------------------------------------------------------

def test_grounded_summary_generator():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    uckr = create_fallback_uckr(norm)
    cfg = OutputGenerationConfig()

    summary = SummaryGenerator.generate(uckr, cfg)
    assert len(summary.text) > 20
    assert len(summary.source_facts) >= 1
    assert all(fid.startswith("F") for fid in summary.source_facts)


def test_grounded_linkedin_generator():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    uckr = create_fallback_uckr(norm)
    cfg = OutputGenerationConfig()

    li = LinkedInGenerator.generate(uckr, cfg)
    assert len(li.headline) > 5
    assert len(li.post_content) > 30
    assert len(li.source_facts) >= 1


def test_grounded_presentation_generator():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    uckr = create_fallback_uckr(norm)
    cfg = OutputGenerationConfig(slide_count=3)

    pres = PresentationGenerator.generate(uckr, cfg)
    assert len(pres.slides) >= 2
    assert pres.slides[0].slide_number == 1
    assert len(pres.source_facts) >= 1
    assert all(isinstance(s.source_facts, list) for s in pres.slides)


def test_grounded_video_generator():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    uckr = create_fallback_uckr(norm)
    cfg = OutputGenerationConfig(video_duration=30)

    video = VideoGenerator.generate(uckr, cfg)
    assert len(video.storyboard) >= 2
    assert len(video.source_facts) >= 1
    assert video.storyboard[0].visual_prompt is not None
    assert video.storyboard[0].narration is not None


def test_grounded_twitter_and_advisory_generators():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="AI in Smart Cities")
    uckr = create_fallback_uckr(norm)
    cfg = OutputGenerationConfig()

    tw = TwitterGenerator.generate(uckr, cfg)
    assert len(tw.thread) >= 2
    assert len(tw.source_facts) >= 1

    adv = AdvisoryGenerator.generate(uckr, cfg)
    assert len(adv.executive_summary) > 10
    assert len(adv.sections) >= 2
    assert len(adv.source_facts) >= 1


# ---------------------------------------------------------------------------
# FastAPI API Endpoint Tests
# ---------------------------------------------------------------------------

def test_api_extract_endpoint_raw_text():
    response = client.post(
        "/consistency/extract",
        json={"raw_text": SAMPLE_ARTICLE_TEXT, "title": "Smart Cities"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["normalized_source"]["source_id"].startswith("SRC-")
    assert "traffic management" in data["normalized_source"]["raw_text"]


def test_api_extract_endpoint_pdf_upload():
    pdf_bytes = create_sample_pdf("AI enhances traffic and energy monitoring.")
    response = client.post(
        "/consistency/extract",
        files={"file": ("smart_cities.pdf", pdf_bytes, "application/pdf")},
        data={"title": "Smart Cities PDF"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["normalized_source"]["source_type"] == "pdf"


def test_api_analyze_endpoint():
    response = client.post(
        "/consistency/analyze",
        json={"raw_text": SAMPLE_ARTICLE_TEXT, "title": "Smart Cities AI"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    uckr = data["uckr"]
    assert len(uckr["facts"]) >= 2
    assert uckr["facts"][0]["id"] == "F001"
    assert len(uckr["key_concepts"]) >= 1


def test_api_registry_facts_endpoint():
    # First analyze to populate registry
    analyze_resp = client.post(
        "/consistency/analyze",
        json={"raw_text": SAMPLE_ARTICLE_TEXT, "title": "Registry Test Doc"}
    )
    source_id = analyze_resp.json()["uckr"]["document"]["id"]

    # Query registry facts
    reg_resp = client.get(f"/consistency/registry/{source_id}/facts")
    assert reg_resp.status_code == 200
    reg_data = reg_resp.json()
    assert reg_data["source_id"] == source_id
    assert reg_data["total_facts"] >= 2
    assert len(reg_data["facts"]) >= 2
    assert reg_data["facts"][0]["id"] == "F001"


def test_api_generate_endpoint():
    norm = extract_from_raw_text(SAMPLE_ARTICLE_TEXT, title="Generate Test Doc")
    uckr = create_fallback_uckr(norm)

    response = client.post(
        "/consistency/generate",
        json={
            "uckr": uckr.model_dump(),
            "config": {
                "output_types": ["summary", "linkedin", "presentation", "video"],
                "audience": "Enterprise Executives",
                "tone": "Strategic",
                "slide_count": 3,
                "video_duration": 20
            }
        }
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert "summary" in data["outputs"]
    assert "linkedin" in data["outputs"]
    assert "presentation" in data["outputs"]
    assert "video" in data["outputs"]
    assert "consistency_audit" in data
    assert data["consistency_audit"]["citation_integrity_score"] == 1.0


def test_api_full_pipeline_endpoint_text():
    response = client.post(
        "/consistency/pipeline",
        data={
            "raw_text": SAMPLE_ARTICLE_TEXT,
            "title": "Smart Cities End-to-End",
            "output_types": "summary,linkedin,presentation,video,twitter,advisory",
            "audience": "Government Leaders",
            "tone": "Policy and Innovation"
        }
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["source"]["source_id"].startswith("SRC-")
    assert len(data["uckr"]["facts"]) >= 2
    assert "summary" in data["outputs"]
    assert "linkedin" in data["outputs"]
    assert "presentation" in data["outputs"]
    assert "video" in data["outputs"]
    assert "twitter" in data["outputs"]
    assert "advisory" in data["outputs"]
    assert data["consistency_audit"]["total_registered_facts"] >= 2
    assert data["consistency_audit"]["fact_coverage_percentage"] > 0.0


def test_api_full_pipeline_endpoint_pdf_upload():
    pdf_bytes = create_sample_pdf("Artificial Intelligence is driving smart traffic routing and sustainable energy grids across 50 global cities.")
    response = client.post(
        "/consistency/pipeline",
        files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
        data={
            "title": "Smart Cities Report",
            "output_types": "summary,linkedin,presentation"
        }
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["source"]["source_type"] == "pdf"
    assert "summary" in data["outputs"]
    assert "linkedin" in data["outputs"]
    assert "presentation" in data["outputs"]
    assert len(data["outputs"]["presentation"]["slides"]) >= 2
    assert data["consistency_audit"]["citation_integrity_score"] == 1.0
