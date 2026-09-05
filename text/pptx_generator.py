import io
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Design System Palette (Executive Modern Tech Theme)
COLOR_DARK_BG = RGBColor(15, 23, 42)        # Slate 900
COLOR_TITLE_DARK = RGBColor(255, 255, 255)  # White
COLOR_SUBTITLE_DARK = RGBColor(148, 163, 184)# Slate 400

COLOR_LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
COLOR_HEADER = RGBColor(30, 41, 59)          # Slate 800
COLOR_ACCENT = RGBColor(37, 99, 235)         # Blue 600
COLOR_TEXT_BODY = RGBColor(51, 65, 85)       # Slate 700
COLOR_MUTED = RGBColor(100, 116, 139)       # Slate 500
COLOR_BOX_BG = RGBColor(255, 255, 255)      # White
COLOR_CARD_BORDER = RGBColor(226, 232, 240)  # Slate 200


def create_pptx_presentation(presentation_data: Dict[str, Any]) -> io.BytesIO:
    """
    Generate a modern, executive 16:9 PowerPoint (.pptx) file from structured presentation JSON.
    Includes custom layouts, bullet points, two-column grids, visual recommendations, and speaker notes.
    """
    prs = Presentation()
    
    # Set 16:9 Widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]  # Blank slide layout
    
    title_text = presentation_data.get("presentation_title") or presentation_data.get("title") or "Executive Presentation"
    subtitle_text = presentation_data.get("subtitle") or presentation_data.get("main_message") or ""
    slides_data = presentation_data.get("slides") or []

    # If no slides array was returned, construct slide from raw contents
    if not slides_data:
        slides_data = [
            {
                "slide_number": 1,
                "title": title_text,
                "layout": "title",
                "subtitle": subtitle_text,
                "speaker_notes": "Welcome audience to the presentation."
            }
        ]

    # --- 1. TITLE SLIDE (Cover Slide) ---
    first_slide_data = slides_data[0]
    cover_slide = prs.slides.add_slide(blank_layout)
    _build_title_slide(cover_slide, title_text, subtitle_text, first_slide_data.get("speaker_notes", ""))

    # --- 2. CONTENT SLIDES ---
    # Process remaining slides (or start from 0 if first slide is content)
    content_slides = slides_data[1:] if first_slide_data.get("layout") == "title" else slides_data
    
    for idx, slide_info in enumerate(content_slides, start=2):
        slide = prs.slides.add_slide(blank_layout)
        layout_type = slide_info.get("layout", "bullet_points").lower()
        
        if "two_column" in layout_type or "grid" in layout_type:
            _build_two_column_slide(slide, slide_info)
        elif "quote" in layout_type or "callout" in layout_type:
            _build_quote_slide(slide, slide_info)
        else:
            _build_bullet_slide(slide, slide_info)
            
        # Attach Speaker Notes
        notes = slide_info.get("speaker_notes", "")
        if notes and hasattr(slide, "notes_slide"):
            slide.notes_slide.notes_text_frame.text = notes

    # Save presentation to BytesIO stream
    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)
    return output_stream


def _build_title_slide(slide, title: str, subtitle: str, speaker_notes: str):
    """Build a striking Dark Cover Title Slide."""
    # Dark Background Rect
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_DARK_BG
    bg.line.fill.background()

    # Accent Header Bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT
    bar.line.fill.background()

    # Title & Subtitle Text Box
    tx_box = slide.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.8), Inches(3.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TITLE_DARK
    p_title.space_after = Pt(14)
    
    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(22)
        p_sub.font.color.rgb = COLOR_SUBTITLE_DARK
        
    if speaker_notes and hasattr(slide, "notes_slide"):
        slide.notes_slide.notes_text_frame.text = speaker_notes


def _add_slide_header(slide, title: str):
    """Add a clean light header banner with blue accent line to a content slide."""
    # Top Header Background
    header_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
    header_bg.fill.solid()
    header_bg.fill.fore_color.rgb = COLOR_LIGHT_BG
    header_bg.line.fill.background()

    # Blue Accent Bar under Header
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()

    # Header Title Text
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.733), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Segoe UI"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_HEADER


def _add_visual_recommendation(slide, visual_text: str):
    """Add an optional visual recommendation footer card at the bottom of the slide."""
    if not visual_text:
        return
        
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.8))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(239, 246, 255)  # Soft blue tint
    card.line.color.rgb = COLOR_ACCENT

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.533), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"💡 Visual Recommendation: {visual_text}"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = COLOR_ACCENT


def _build_bullet_slide(slide, slide_info: Dict[str, Any]):
    """Build a Standard Bullet Points Slide."""
    title = slide_info.get("title", "Key Highlights")
    _add_slide_header(slide, title)
    
    # Bullet Content Frame
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    content = slide_info.get("content", [])
    if isinstance(content, str):
        content = [content]

    for idx, item in enumerate(content):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(14)
        
    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""))


def _build_two_column_slide(slide, slide_info: Dict[str, Any]):
    """Build a Two-Column Grid Slide."""
    title = slide_info.get("title", "Comparative Insights")
    _add_slide_header(slide, title)

    col_left = slide_info.get("column_left") or slide_info.get("content", [])[:3]
    col_right = slide_info.get("column_right") or slide_info.get("content", [])[3:]

    if isinstance(col_left, str):
        col_left = [col_left]
    if isinstance(col_right, str):
        col_right = [col_right]

    # Column Left Box
    card_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.4))
    card_left.fill.solid()
    card_left.fill.fore_color.rgb = COLOR_BOX_BG
    card_left.line.color.rgb = COLOR_CARD_BORDER

    tb_l = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    for idx, item in enumerate(col_left):
        p = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(12)

    # Column Right Box
    card_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(1.6), Inches(5.6), Inches(4.4))
    card_right.fill.solid()
    card_right.fill.fore_color.rgb = COLOR_BOX_BG
    card_right.line.color.rgb = COLOR_CARD_BORDER

    tb_r = slide.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(4.0))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    for idx, item in enumerate(col_right):
        p = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(12)

    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""))


def _build_quote_slide(slide, slide_info: Dict[str, Any]):
    """Build a Modern Callout / Quote Slide."""
    title = slide_info.get("title", "Key Takeaway")
    _add_slide_header(slide, title)

    quote_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.0), Inches(10.333), Inches(3.6))
    quote_card.fill.solid()
    quote_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
    quote_card.line.color.rgb = COLOR_ACCENT

    tb = slide.shapes.add_textbox(Inches(1.8), Inches(2.3), Inches(9.733), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    quote_text = slide_info.get("content")
    if isinstance(quote_text, list):
        quote_text = " ".join(quote_text)

    p = tf.paragraphs[0]
    p.text = f'"{quote_text}"'
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.italic = True
    p.font.color.rgb = COLOR_HEADER
    p.alignment = PP_ALIGN.CENTER

    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""))
